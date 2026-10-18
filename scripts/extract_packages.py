#!/usr/bin/env python3
"""Stage D, first half: the four Information Package workbooks to an intermediate tree.

The library says what a class and an attribute mean. These files say how they are assembled into a
message: the information path, how many times each node may occur, and which derived packages - the
real messages, eTIR, Certificate of Origin, Maritime API and the rest - use it.

Columns are resolved by their heading rather than by position, because the four files do not share a
layout: MDIP has no Remark column, so its Min and Max sit one place to the left of everyone else's.

    python3 scripts/extract_packages.py
"""
import json
import pathlib
import sys

import openpyxl

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import _release  # noqa: E402

ROOT = pathlib.Path(__file__).resolve().parent.parent

# key -> (sheet, human name, what it is). The file is whichever workbook of the release is that
# package; releases have renamed them before.
PACKAGES = {
    "mdip": ("MDIP (MetaData)", "Meta Data Information Package",
             "The technical aspect of a data exchange, kept apart so the other packages carry only business content."),
    "aip": ("AIP (RES and InterGov)", "Additional Information Package",
            "A wrapper around business information: Response for government to business, Intergov for government to government."),
    "bip-dec": ("BIP DEC", "Declaration Base Information Package",
                "Business to government submissions: import and export declarations, cargo reports, transit declarations."),
    "bip-lpco": ("BIP LPCO", "LPCO Base Information Package",
                 "Licences, permits, certificates and other documents."),
}

# Each workbook groups its columns with two heading rows above the column names. The groups that
# hold derived packages are named in the second row; a group named in the first row is a detail
# block - the IMO Compendium mapping, the per-DIP names and definitions - and is not a flag.
# Matching on the group rather than on the column is what keeps the Declaration's four derived
# packages from becoming forty.
# 4.0.0 and 4.1.0 name the same blocks differently, and the Declaration's "Main Procedures" block
# (IM, EX, CRI, CRE, CONV, TRT) is deliberately absent: those flag which customs procedure uses a
# position, which is a different axis from which message does.
DERIVED_GROUP = ("derived ip", "dip use", "intergov projects", "intergov dip",
                 "intergov work in progress", "response category", "response dip",
                 "ig projects", "ig derived ip", "res category", "res derived ip")


def column_groups(row1, row2):
    """Column index -> (group heading, which row named it). A column belongs to the nearest group
    heading at or to its left."""
    starts = {}
    for level, row in (("row1", row1), ("row2", row2)):
        for i, cell in enumerate(row):
            if cell and str(cell).strip():
                starts[i] = (norm(cell), level)
    out, current = {}, (None, None)
    for i in range(max(len(row1), len(row2))):
        if i in starts:
            current = starts[i]
        out[i] = current
    return out


def norm(text):
    return " ".join(str(text or "").replace("_x000D_", " ").split()).strip().lower()


def read_package(key, release, version):
    sheet, title, description = PACKAGES[key]
    filename = _release.workbook(key, release).name
    wb = openpyxl.load_workbook(release / filename, read_only=True, data_only=True)
    sheet = sheet if sheet in wb.sheetnames else next(
        (s for s in wb.sheetnames if s.lower().startswith(sheet.split(" ")[0].lower())),
        wb.sheetnames[0])
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    headers = [norm(c) for c in rows[2]]
    col = {h: i for i, h in enumerate(headers) if h}
    groups = column_groups(rows[0], rows[1])

    def get(row, *names):
        for n in names:
            if n in col and row[col[n]] not in (None, ""):
                return str(row[col[n]]).strip()
        return None

    path_col = next((col[h] for h in ("uml model path", "xpath") if h in col), 0)
    body = [r for r in rows[3:] if r[path_col] and str(r[path_col]).strip()]
    # A derived package is a column whose nearest group heading was named in the second row and is
    # one of the derived blocks. The detail blocks that follow - the IMO Compendium mapping, the
    # per-certificate box numbers in 4.0.0 - are named in the first row, which is what ends the
    # span without a column count having to be written down anywhere.
    dips, dip_group, dip_col = [], {}, {}
    for i, h in enumerate(headers):
        if not h:
            continue
        group, level = groups.get(i, (None, None))
        if level != "row2" or not group or not group.startswith(DERIVED_GROUP):
            continue
        if sum(1 for r in body if i < len(r) and r[i] not in (None, "")):
            dips.append(h)
            dip_group[h] = group
            dip_col[h] = i

    nodes = []
    for n, r in enumerate(body, start=4):
        path = str(r[path_col]).strip()
        node = {
            "path": path,
            "depth": len(path.split("/")),
            "wco_id": get(r, "wco id"),
            "name": get(r, "wco name"),
            "definition": get(r, "wco definition"),
            "remark": get(r, "remark"),
            "format": get(r, "wco format"),
            "min": get(r, "min"),
            "max": get(r, "max"),
            "code_list": get(r, "wco code list"),
            "xml_tag": get(r, "xml tag"),
            "json_tag": get(r, "json tag"),
            "type": get(r, "type"),            # C for a class, E for an element
            "seq": get(r, "seq. no."),
            "uci": get(r, "dec uci", "lpco uci", "mdip uci", "ig uci", "res uci",
                       "ig", "res", "class"),
            "used_by": sorted(h for h in dips
                              if dip_col[h] < len(r) and r[dip_col[h]] not in (None, "")),
            "source": {"file": filename, "sheet": sheet, "row": n},
        }
        nodes.append(node)

    wb.close()
    return {
        "key": key, "title": title, "description": description,
        "source": {"file": filename, "sheet": sheet, "version": version},
        "derived_packages": dips,
        "derived_package_groups": dip_group,
        "nodes": nodes,
    }


def extract(release):
    version = _release.version(release)
    out = ROOT / "intermediate" / f"v{version}" / "packages"
    out.mkdir(parents=True, exist_ok=True)
    summary = {}
    for key in PACKAGES:
        pkg = read_package(key, release, version)
        (out / f"{key}.json").write_text(json.dumps(pkg, indent=1, ensure_ascii=False) + "\n")
        classes = sum(1 for n in pkg["nodes"] if n["type"] == "C")
        elements = sum(1 for n in pkg["nodes"] if n["type"] == "E")
        repeating = sum(1 for n in pkg["nodes"] if (n["max"] or "").lower() == "unbounded")
        summary[key] = {
            "title": pkg["title"], "nodes": len(pkg["nodes"]),
            "classes": classes, "elements": elements, "repeating": repeating,
            "max_depth": max(n["depth"] for n in pkg["nodes"]),
            "derived_packages": pkg["derived_packages"],
        }
    (out / "summary.json").write_text(json.dumps(summary, indent=1, ensure_ascii=False) + "\n")
    total = sum(s["nodes"] for s in summary.values())
    dips = sum(len(s["derived_packages"]) for s in summary.values())
    print(f"  {version}: {total} positions across {len(summary)} packages, {dips} derived packages")
    for key, s in summary.items():
        print(f"      {key:9} {s['nodes']:5} positions, {len(s['derived_packages'])} derived: "
              f"{', '.join(s['derived_packages']) or 'none'}")
    return summary


def main():
    releases = _release.releases() if "--all" in sys.argv else [_release.release_dir()]
    for release in releases:
        extract(release)
    return 0


if __name__ == "__main__":
    sys.exit(main())
